"""File loaders: turn PDF / DOCX / TXT / images into a flat list of text items and image items.

Mixed documents (e.g. a PDF with text AND pictures/charts/scans) produce both kinds of items,
tagged with the same source + page so they can be linked again at retrieval time.
"""
import io
import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from PIL import Image

from . import config

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".json", ".log"}
SUPPORTED = IMAGE_EXT | TEXT_EXT | {".pdf", ".docx"}


@dataclass
class RawItem:
    type: str                    # "text" | "image"
    source: str
    page: int | None
    text: str = ""
    image: bytes | None = None   # normalised JPEG bytes when type == "image"


# ----------------------------------------------------------------------------- images
def normalize_image(data: bytes, min_side: int | None = None) -> bytes | None:
    """Decode -> RGB -> downscale -> JPEG that fits the API size budget. None if unusable/too small."""
    min_side = config.MIN_IMAGE_SIDE if min_side is None else min_side
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:
        return None
    if min(img.size) < min_side:
        return None

    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        bg = Image.new("RGB", rgba.size, "white")
        bg.paste(rgba, mask=rgba.split()[3])
        img = bg
    else:
        img = img.convert("RGB")

    longest = max(img.size)
    if longest > config.MAX_IMAGE_SIDE:
        s = config.MAX_IMAGE_SIDE / longest
        img = img.resize((max(1, int(img.width * s)), max(1, int(img.height * s))), Image.LANCZOS)

    quality = 85
    while True:
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        if buf.tell() <= config.MAX_IMAGE_BYTES or min(img.size) <= 64:
            return buf.getvalue()
        if quality > 55:
            quality -= 10
        else:
            img = img.resize((int(img.width * 0.85), int(img.height * 0.85)), Image.LANCZOS)


# ----------------------------------------------------------------------------- loaders
def _load_pdf(path: Path, source: str) -> list[RawItem]:
    items: list[RawItem] = []
    doc = pymupdf.open(path)
    seen_xrefs: set[int] = set()
    try:
        for pno, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                items.append(RawItem("text", source, pno, text=text))

            embedded = 0
            for info in page.get_images(full=True):
                xref = info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                data = None
                try:
                    pix = pymupdf.Pixmap(doc, xref)
                    if pix.n - pix.alpha >= 4:          # CMYK -> RGB
                        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                    data = normalize_image(pix.tobytes("png"))
                except Exception:
                    try:
                        data = normalize_image(doc.extract_image(xref)["image"])
                    except Exception:
                        data = None
                if data:
                    embedded += 1
                    items.append(RawItem("image", source, pno, image=data))

            # Scanned page (no text layer, no usable embedded image) or chart-heavy vector page:
            # render the whole page so the VLM can read it.
            try:
                drawings = len(page.get_drawings())
            except Exception:
                drawings = 0
            scanned = len(text) < config.SCANNED_PAGE_MIN_CHARS and embedded == 0
            vector_heavy = drawings >= config.VECTOR_PAGE_DRAWINGS
            if scanned or vector_heavy:
                data = normalize_image(page.get_pixmap(dpi=config.PAGE_RENDER_DPI).tobytes("png"), min_side=0)
                if data:
                    items.append(RawItem("image", source, pno, image=data))
    finally:
        doc.close()
    return items


def _load_docx(path: Path, source: str) -> list[RawItem]:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    items: list[RawItem] = []
    if parts:
        items.append(RawItem("text", source, None, text="\n".join(parts)))
    for rel in d.part.rels.values():
        if rel.is_external or "/image" not in rel.reltype:
            continue
        data = normalize_image(rel.target_part.blob)
        if data:
            items.append(RawItem("image", source, None, image=data))
    return items


def load_file(path: str | Path, source: str | None = None) -> list[RawItem]:
    path = Path(path)
    source = source or path.name
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path, source)
    if ext == ".docx":
        return _load_docx(path, source)
    if ext in IMAGE_EXT:
        data = normalize_image(path.read_bytes(), min_side=0)
        return [RawItem("image", source, 1, image=data)] if data else []
    if ext in TEXT_EXT:
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        return [RawItem("text", source, None, text=text)] if text else []
    raise ValueError(f"Unsupported file type: {ext or '(none)'}. Supported: {', '.join(sorted(SUPPORTED))}")


# ----------------------------------------------------------------------------- chunking
def chunk_text(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    size = size or config.CHUNK_SIZE
    overlap = config.CHUNK_OVERLAP if overlap is None else overlap
    overlap = min(overlap, size // 2)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    chunks: list[str] = []
    cur = ""
    for para in text.split("\n\n"):
        if len(para) > size:                       # hard-split very long paragraphs
            if cur:
                chunks.append(cur)
                cur = ""
            step = max(1, size - overlap)
            chunks.extend(para[i:i + size] for i in range(0, len(para), step))
            continue
        if len(cur) + len(para) + 2 <= size:
            cur = f"{cur}\n\n{para}" if cur else para
        else:
            chunks.append(cur)
            cur = (cur[-overlap:] + "\n\n" + para) if overlap and cur else para
    if cur:
        chunks.append(cur)
    return [c.strip() for c in chunks if c.strip()]
