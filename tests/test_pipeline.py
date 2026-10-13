import pymupdf
import pytest
from fastapi.testclient import TestClient

from mmrag import api, config
from mmrag.ingest import chunk_text, load_file
from tests.conftest import make_png


def make_mixed_pdf(path):
    """Page 1: text + embedded chart image. Page 2: text only. Page 3: scanned-style (image only)."""
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "Annual report. The company was founded in 2019 in Delhi and builds AI agents.", fontsize=11)
    p1.insert_image(pymupdf.Rect(72, 120, 472, 420), stream=make_png())
    p2 = doc.new_page()
    p2.insert_text((72, 72), "Page two only has text about the hiring plan for Germany and Sweden.", fontsize=11)
    p3 = doc.new_page()
    p3.insert_image(pymupdf.Rect(50, 50, 550, 700), stream=make_png(500, 650))
    doc.save(path)
    doc.close()


def test_pdf_yields_text_and_image_items(tmp_path):
    pdf = tmp_path / "mixed.pdf"
    make_mixed_pdf(pdf)
    items = load_file(pdf)
    kinds = [(i.type, i.page) for i in items]
    assert ("text", 1) in kinds and ("image", 1) in kinds     # mixed page
    assert ("text", 2) in kinds                                # text-only page
    assert ("image", 3) in kinds and ("text", 3) not in kinds  # scanned/image-only page


def test_mixed_pdf_query_returns_both_modalities(rag):
    import tempfile, os
    with tempfile.TemporaryDirectory() as d:
        pdf = os.path.join(d, "mixed.pdf")
        make_mixed_pdf(pdf)
        stats = rag.ingest(pdf, "mixed.pdf")
    assert stats["text_chunks"] >= 2 and stats["image_chunks"] >= 2 and stats["images_failed"] == 0

    # k=1 -> only the chart description matches directly; page expansion must add page-1 text too
    res = rag.query("what does the revenue chart show for Q3?", k=1)
    types = {(s["page"], s["type"]) for s in res["sources"]}
    assert (1, "image") in types
    assert (1, "text") in types, "same-page text should be pulled in next to the matching image"
    assert rag.client.last_images, "original image should be attached for the VLM"

    res = rag.query("hiring plan Germany Sweden", k=1)
    assert any(s["page"] == 2 and s["type"] == "text" for s in res["sources"])


def test_image_only_document(rag, tmp_path):
    png = tmp_path / "scan.png"
    png.write_bytes(make_png(300, 200))
    stats = rag.ingest(png)
    assert stats["image_chunks"] >= 1 and stats["text_chunks"] == 0
    res = rag.query("quarterly revenue chart")
    assert res["sources"][0]["type"] == "image"


def test_text_only_document(rag, tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("Arbaz builds LangGraph agents.\n\nHe also uses FastMCP and RAG pipelines.")
    stats = rag.ingest(f)
    assert stats["text_chunks"] >= 1 and stats["image_chunks"] == 0
    res = rag.query("which agents framework")
    assert res["sources"][0]["type"] == "text"
    assert rag.client.last_images == []


def test_docx_text_and_image(rag, tmp_path):
    import docx
    img = tmp_path / "chart.png"
    img.write_bytes(make_png())
    d = docx.Document()
    d.add_paragraph("Quarterly summary written inside a Word document.")
    d.add_picture(str(img))
    path = tmp_path / "report.docx"
    d.save(path)
    stats = rag.ingest(path)
    assert stats["text_chunks"] >= 1 and stats["image_chunks"] >= 1


def test_reingest_replaces_and_delete_cleans_up(rag, tmp_path):
    png = tmp_path / "scan.png"
    png.write_bytes(make_png())
    rag.ingest(png)
    rag.ingest(png)                                           # same file again
    assert rag.documents()["scan.png"]["image_chunks"] == 1   # no duplicates
    assert len(list(config.IMAGE_DIR.glob("*.jpg"))) == 1
    assert rag.delete("scan.png") >= 1
    assert rag.documents() == {} and list(config.IMAGE_DIR.glob("*.jpg")) == []


def test_unsupported_type(rag, tmp_path):
    f = tmp_path / "x.xyz"
    f.write_text("hi")
    with pytest.raises(ValueError):
        rag.ingest(f)


def test_chunker_respects_size():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 60 for i in range(12))
    chunks = chunk_text(text, size=500, overlap=50)
    assert len(chunks) > 1 and all(len(c) <= 620 for c in chunks)


def test_api_end_to_end(rag, tmp_path):
    api._rag = rag
    client = TestClient(api.app)
    assert client.get("/health").status_code == 200
    r = client.post("/ingest", files=[("files", ("a.txt", b"Delhi is the capital of India.", "text/plain")),
                                      ("files", ("b.png", make_png(), "image/png"))])
    assert r.status_code == 200 and all("error" not in x for x in r.json()["results"])
    q = client.post("/query", json={"question": "capital of India?"})
    assert q.status_code == 200 and q.json()["sources"]
    assert client.delete("/documents/a.txt").status_code == 200
    assert client.delete("/documents/a.txt").status_code == 404
    api._rag = None
